import collections
import collections.abc
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
import os

prs = Presentation('C:/Users/msı/Desktop/wsn_api/stockfish_sunumum.pptx')

# Slide layout 1 (usually Title and Content)
bullet_slide_layout = prs.slide_layouts[1]

# Slide 1: Yeni Geliştirmeler ve Mimari
slide1 = prs.slides.add_slide(bullet_slide_layout)
title_shape1 = slide1.shapes.title
body_shape1 = slide1.shapes.placeholders[1]

title_shape1.text = "Sistem Güncellemeleri ve Yeni Mimari"

tf1 = body_shape1.text_frame
tf1.text = "Genişletilmiş Ağ Desteği: 10x10 WSN ızgaralarını desteklemek için Fairy-Stockfish Largeboard entegrasyonu sağlandı."
p = tf1.add_paragraph()
p.text = "Sunucusuz (Serverless) Altyapı: Backend Vercel Serverless üzerine taşınarak soğuk başlama (cold start) süresi ve yanıt gecikmesi minimuma indirildi."
p.level = 0
p = tf1.add_paragraph()
p.text = "Dinamik Varyant Yönetimi: Ağ topolojisi, özel kurallar ve boyutlar variants.ini dosyasından anlık olarak parse edilerek sisteme tanıtıldı."
p.level = 0
p = tf1.add_paragraph()
p.text = "Çift Sunucu Yedekliliği (Fallback): Vercel ve Render uç noktaları aynı anda kullanılarak kesintisiz API iletişimi sağlandı."
p.level = 0

# Slide 2: Arayüz ve Görselleştirme
slide2 = prs.slides.add_slide(bullet_slide_layout)
title_shape2 = slide2.shapes.title
body_shape2 = slide2.shapes.placeholders[1]

title_shape2.text = "Arayüz (UI) ve Görselleştirme"

tf2 = body_shape2.text_frame
tf2.text = "Dinamik WSN Izgarası: FEN dizilimine göre düğümlerin (sensörlerin) otomatik konumlandırıldığı modern görsel arayüz."
p = tf2.add_paragraph()
p.text = "Karar Filtreleme Paneli: Gelen CP skorlarının belirli eşik değerleriyle (threshold) test edildiği otonom karar onayı/reddi sistemi."
p.level = 0
p = tf2.add_paragraph()
p.text = "Çoklu Dil (i18n): Sistem arayüzünün Türkçe ve İngilizce olarak dinamik JSON dosyalarıyla yönetilmesi."
p.level = 0
p = tf2.add_paragraph()
p.text = "Hata Yönetimi: Ağ kopmalarına karşı zaman aşımı (timeout) kontrolü ve kullanıcıya anlık bildirim sistemi."
p.level = 0

prs.save('C:/Users/msı/Desktop/wsn_api/stockfish_sunumum_guncel.pptx')
print("Sunum guncellendi!")
